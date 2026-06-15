from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Bao_cao_quy_trinh_AI_VibeCheck.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8F1FB"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "6B7280"
BORDER = "DADCE0"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_font(paragraph, size=11, color=None, bold=False):
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    title.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("VibeCheck AI Repository Analysis Report")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)


def add_title_page(document):
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Báo cáo quy trình phân tích Code Smell và Bảo mật bằng AI")

    subtitle = document.add_paragraph()
    subtitle.add_run(
        "Hệ thống VibeCheck - các mục 4.2 đến 4.6"
    ).bold = True
    set_paragraph_font(subtitle, size=13, color=MID_GRAY)

    meta = document.add_paragraph()
    meta.add_run("Phạm vi báo cáo: ").bold = True
    meta.add_run(
        "detection phase, validation phase, security workflow, kiến trúc hệ thống, scoring/severity và cơ chế Generate Prompt."
    )
    set_paragraph_font(meta)

    note = document.add_paragraph()
    note.add_run("Cơ sở kỹ thuật: ").bold = True
    note.add_run(
        "báo cáo được tổng hợp từ các service backend hiện tại của VibeCheck, bao gồm pipeline GitHub ingestion, AI provider integration, cache theo file SHA, các rule file trong thư mục docs và logic tạo report."
    )
    set_paragraph_font(note)

    document.add_paragraph()
    callout = document.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_cell_margins(callout, top=120, bottom=120, start=160, end=160)
    cell = callout.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_border(cell, color="BFD7EE")
    p = cell.paragraphs[0]
    p.add_run("Tóm tắt định hướng: ").bold = True
    p.add_run(
        "VibeCheck không yêu cầu AI đọc toàn bộ repository một lần. Hệ thống lọc file, ưu tiên file quan trọng, chia batch, chạy detection song song có giới hạn, cache kết quả theo GitHub blob SHA, sau đó validation một lần để chuẩn hóa findings trước khi lưu DB và trả về FE."
    )
    set_paragraph_font(p)

    document.add_page_break()


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)
        set_paragraph_font(p)


def add_numbered(document, items):
    for item in items:
        p = document.add_paragraph(style="List Number")
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)
        set_paragraph_font(p)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_cell_margins(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)

    for idx, text in enumerate(headers):
        cell = header_row.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.add_run(text).bold = True
        set_paragraph_font(p, size=10, color=DARK_BLUE, bold=True)
        if widths:
            cell.width = Inches(widths[idx])

    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.add_run(str(text))
            set_paragraph_font(p, size=10)
            if widths:
                cell.width = Inches(widths[idx])

    document.add_paragraph()
    return table


def add_section_42(document):
    document.add_heading("4.2. Quy trình phát hiện code smell và bảo mật bằng AI", level=1)
    p = document.add_paragraph()
    p.add_run(
        "Quy trình phân tích của VibeCheck được thiết kế theo mô hình hai pha: Detection Phase và Validation Phase. "
    )
    p.add_run(
        "Pha đầu tiên ưu tiên phát hiện đầy đủ các dấu hiệu nghi vấn, còn pha thứ hai ưu tiên lọc nhiễu và chuẩn hóa kết quả để hiển thị cho người dùng."
    )
    set_paragraph_font(p)

    add_table(
        document,
        ["Pha", "Mục tiêu", "Đầu vào", "Đầu ra"],
        [
            [
                "Detection Phase",
                "Tăng độ bao phủ, phát hiện càng nhiều dấu hiệu nghi vấn càng tốt.",
                "Nội dung từng file hoặc batch file, Rule 1, prompt detection.",
                "Danh sách raw detections có evidence, metrics, severity và suggestion ban đầu.",
            ],
            [
                "Validation Phase",
                "Tăng độ chính xác, loại false positive và chuẩn hóa dữ liệu.",
                "Raw detections từ phase 1, Rule 2, prompt validation.",
                "Danh sách findings cuối cùng đã deduplicate, có severityDisplay và suggestion cụ thể.",
            ],
        ],
        widths=[1.45, 2.0, 1.65, 1.75],
    )

    document.add_heading("4.2.1. Detection Phase", level=2)
    p = document.add_paragraph()
    p.add_run("Prompt 1. ").bold = True
    p.add_run(
        "Prompt detection yêu cầu AI đóng vai senior software engineer chuyên static analysis. Với code smell, prompt buộc AI bám theo CODE_SMELL_RULE_1.md; với security, prompt buộc AI bám theo SECURITY_RULE_1.md. AI phải trả về JSON dạng raw list, chưa phải response cuối cùng."
    )
    set_paragraph_font(p)

    add_bullets(
        document,
        [
            (
                "Rule 1: ",
                "định nghĩa các nhóm code smell như Bloaters, Change Preventers, Couplers, Dispensables, OO Abusers; đồng thời quy định metrics như method_length, class_size, nesting_depth, external_dependency_ratio và duplication_signals.",
            ),
            (
                "High recall strategy: ",
                "AI được hướng dẫn inclusive, chấp nhận các smell còn chưa chắc chắn và không lọc quá mạnh ở pha này. Mục tiêu là giảm rủi ro bỏ sót vấn đề.",
            ),
            (
                "Evidence-first output: ",
                "mỗi detection phải có filePath, title, description, evidence, line range, severity và metrics để phục vụ validation sau đó.",
            ),
            (
                "Suggestion ban đầu: ",
                "VibeCheck yêu cầu AI sinh suggestion cụ thể theo issue, không dùng câu mẫu chung cho mọi file.",
            ),
        ],
    )

    p = document.add_paragraph()
    p.add_run("Cách hệ thống thực thi. ").bold = True
    p.add_run(
        "Backend chia các file đã chọn thành batch nhỏ theo ANALYSIS_DETECTION_BATCH_SIZE. Các batch được gửi tới AI provider thông qua callProviderJson và chạy qua concurrency limiter dùng chung, giúp tăng tốc nhưng vẫn tránh vượt ngưỡng request đồng thời."
    )
    set_paragraph_font(p)

    document.add_heading("4.2.2. Validation Phase", level=2)
    p = document.add_paragraph()
    p.add_run("Prompt 2. ").bold = True
    p.add_run(
        "Prompt validation yêu cầu AI đóng vai chuyên gia software quality hoặc application security validation. Input của prompt là toàn bộ raw detections từ phase 1. AI phải đánh giá lại evidence, metrics và rule matching trước khi trả về final JSON."
    )
    set_paragraph_font(p)

    add_bullets(
        document,
        [
            (
                "Rule 2: ",
                "quy định chỉ giữ finding có bằng chứng rõ ràng, reasoning nhất quán, metrics hỗ trợ đúng và không vi phạm định nghĩa rule.",
            ),
            (
                "Precision filtering: ",
                "các detection mơ hồ, trùng lặp, thiếu evidence hoặc confidence thấp phải bị loại bỏ. Với security, issue không có exploit path thực tế cũng bị loại.",
            ),
            (
                "Deduplication and merge: ",
                "các issue cùng loại, cùng vùng code được gộp để report không bị nhiễu.",
            ),
            (
                "Final JSON: ",
                "kết quả cuối cùng gồm index, fileName, filePath, title, description, suggestion, startLine, endLine, severity và severityDisplay.",
            ),
        ],
    )


def add_section_43(document):
    document.add_heading("4.3. Quy trình phân tích bảo mật bằng AI", level=1)
    p = document.add_paragraph()
    p.add_run(
        "Phân tích bảo mật trong VibeCheck dùng cùng kiến trúc hai pha nhưng có bộ rule, metrics và tiêu chí validation riêng. Trọng tâm không chỉ là phát hiện pattern xấu, mà còn đánh giá xem pattern đó có tạo ra đường khai thác thực tế hay không."
    )
    set_paragraph_font(p)

    document.add_heading("OWASP-based workflow", level=2)
    p = document.add_paragraph()
    p.add_run(
        "Workflow bảo mật được xây dựng theo các nhóm rủi ro phổ biến trong OWASP như injection, XSS, authentication/authorization, secrets exposure, cryptography, SSRF, path traversal và insecure configuration. Các nhóm này được mô tả trong SECURITY_RULE_1.md và được dùng làm khung phát hiện ban đầu."
    )
    set_paragraph_font(p)

    add_numbered(
        document,
        [
            "Lọc file có giá trị bảo mật cao như routes, controllers, services, middleware, auth, API layer, Dockerfile, docker-compose, .env.example và .github/workflows.",
            "Chạy detection prompt để tìm dấu hiệu injection, XSS, thiếu authz, hardcoded secret, unsafe fetch/file handling, cấu hình CORS hoặc security header chưa phù hợp.",
            "Thu thập metrics bảo mật như tainted_input_flow, auth_boundary, validation_strength, dangerous_api_usage và sensitive_data_exposure.",
            "Chạy validation prompt để loại issue thiếu exploit path, issue chỉ nằm trong test code hoặc issue chỉ là best practice không có tác động bảo mật rõ ràng.",
            "Chuẩn hóa findings rồi lưu vào collection issue với type là security.",
        ],
    )

    document.add_heading("Security prompting", level=2)
    p = document.add_paragraph()
    p.add_run(
        "Prompt security yêu cầu AI đóng vai senior application security engineer. Mỗi vulnerability phải có category, evidence, severity và remediation suggestion. Suggestion phải chỉ ra input/boundary/API nguy hiểm cụ thể, ví dụ validate query parameter nào, thay thế API nào, hoặc bổ sung ownership check ở route nào."
    )
    set_paragraph_font(p)

    document.add_heading("Vulnerability validation", level=2)
    p = document.add_paragraph()
    p.add_run(
        "Ở validation phase, VibeCheck ưu tiên giữ ít finding hơn nhưng có chất lượng cao hơn. Một issue chỉ được giữ nếu evidence đủ rõ, category khớp với Rule 1 và có realistic exploit path hoặc weakness quan sát được. Những cảnh báo chỉ dựa vào suy đoán hoặc không liên quan production behavior sẽ bị loại."
    )
    set_paragraph_font(p)


def add_section_44(document):
    document.add_heading("4.4. Kiến trúc hệ thống VibeCheck", level=1)
    p = document.add_paragraph()
    p.add_run(
        "VibeCheck được tổ chức theo kiến trúc client-server. Frontend phụ trách trải nghiệm người dùng và hiển thị report; backend phụ trách xác thực, lấy repository, phân loại kích thước, gọi AI, lưu lịch sử và sinh prompt sửa lỗi."
    )
    set_paragraph_font(p)

    add_table(
        document,
        ["Thành phần", "Vai trò chính", "Ghi chú kỹ thuật"],
        [
            [
                "Frontend",
                "Đăng nhập GitHub, nhập repo URL, gọi check/analyze/history, hiển thị findings, generate và regenerate prompt.",
                "FE sử dụng accessToken JWT từ backend sau OAuth callback.",
            ],
            [
                "Backend",
                "Cung cấp REST API, xử lý auth middleware, repository check, AI analysis, history, delete report và prompt generation.",
                "Express, MongoDB, JWT, CORS, rate limit và service layer tách riêng.",
            ],
            [
                "AI provider integration",
                "Dùng Gemini mặc định của hệ thống; nếu repo large hoặc người dùng muốn, dùng OpenAI API key cá nhân.",
                "Provider được chuẩn hóa qua callProviderJson, response phải là JSON.",
            ],
            [
                "Repository ingestion pipeline",
                "Gọi GitHub API, lấy tree, bỏ qua file rác, ưu tiên file quan trọng, tải raw content và tạo code sample.",
                "Có GITHUB_TOKEN để tăng quota GitHub API khi chạy production.",
            ],
            [
                "MongoDB",
                "Lưu users, repositories, roast runs, roast issues và analysis_file_caches.",
                "Cache theo repositoryFullName, filePath, fileSha, provider, model và rulesHash.",
            ],
        ],
        widths=[1.55, 2.85, 2.25],
    )

    document.add_heading("Repository ingestion pipeline", level=2)
    add_numbered(
        document,
        [
            "Người dùng nhập GitHub public repository URL.",
            "Backend parse owner/name và gọi GitHub REST API để lấy repository metadata, default branch và recursive tree.",
            "Hệ thống loại bỏ node_modules, .git, dist, build, coverage, lock file, file media/binary, file minified và các hidden folder không cần phân tích.",
            "Một số file đặc biệt vẫn được giữ vì có giá trị bảo mật: .github/workflows, Dockerfile, docker-compose.yml và .env.example.",
            "File được chấm priority theo path và extension. Các path như src/pages, src/components, services, controllers, routes, middleware, auth, utils được ưu tiên cao.",
            "Backend tải raw content của tối đa ANALYSIS_MAX_PRIORITY_FILES file, mỗi file không vượt ANALYSIS_MAX_FILE_BYTES.",
            "AI detection dùng cache theo GitHub blob SHA. Nếu file không đổi, hệ thống dùng lại raw detections cũ thay vì gọi AI lại.",
        ],
    )

    document.add_heading("AI provider integration", level=2)
    p = document.add_paragraph()
    p.add_run(
        "Với repository small hoặc medium, hệ thống mặc định dùng Gemini thông qua GEMINI_ANALYSIS_MODEL. Với repository large, người dùng bắt buộc nhập OpenAI API key cá nhân; ngoài ra với small/medium, nếu người dùng nhập key thì hệ thống vẫn có thể dùng provider user_openai. Provider và model là một phần của cache key để tránh trộn kết quả giữa Gemini và OpenAI."
    )
    set_paragraph_font(p)


def add_section_45(document):
    document.add_heading("4.5. Cơ chế scoring và severity classification", level=1)
    p = document.add_paragraph()
    p.add_run(
        "Sau validation, mỗi finding được chuẩn hóa severity về ba mức low, medium và high. Severity phản ánh mức độ ảnh hưởng của issue tới maintainability hoặc security. Hệ thống sau đó dùng severity để tính điểm tổng thể và xếp hạng repository."
    )
    set_paragraph_font(p)

    add_table(
        document,
        ["Severity", "Ý nghĩa với code smell", "Ý nghĩa với security", "Trọng số"],
        [
            [
                "Low",
                "Ảnh hưởng nhỏ tới readability hoặc maintainability.",
                "Rủi ro thấp, exploitability chưa rõ hoặc impact hạn chế.",
                "1",
            ],
            [
                "Medium",
                "Ảnh hưởng rõ tới khả năng bảo trì, cần refactor có kiểm soát.",
                "Có weakness thực tế nhưng impact hoặc điều kiện khai thác ở mức trung bình.",
                "2",
            ],
            [
                "High",
                "Ảnh hưởng lớn tới thiết kế, dễ gây lỗi khi mở rộng hoặc sửa đổi.",
                "Có đường khai thác hoặc tác động bảo mật đáng kể.",
                "3",
            ],
        ],
        widths=[0.9, 2.15, 2.55, 0.75],
    )

    p = document.add_paragraph()
    p.add_run("Công thức điểm. ").bold = True
    p.add_run(
        "VibeCheck bắt đầu từ 100 điểm. Mỗi code smell bị trừ severityWeight x 3. Mỗi security issue bị trừ severityWeight x 7. Security có hệ số phạt cao hơn vì lỗi bảo mật thường có tác động rủi ro lớn hơn lỗi maintainability."
    )
    set_paragraph_font(p)

    add_table(
        document,
        ["Khoảng điểm", "Grade"],
        [
            ["90 - 100", "A"],
            ["80 - 89", "B+"],
            ["70 - 79", "B"],
            ["60 - 69", "C+"],
            ["50 - 59", "C"],
            ["40 - 49", "D"],
            ["0 - 39", "F"],
        ],
        widths=[2.0, 1.0],
    )

    p = document.add_paragraph()
    p.add_run("Repo info. ").bold = True
    p.add_run(
        "Report trả về số file, size label, tổng code smell, tổng security, tổng high, tổng medium, tổng low và analysisType. Riêng từng nhóm codeSmells và security có all, totalHigh, totalMedium, totalLow và danh sách items."
    )
    set_paragraph_font(p)


def add_section_46(document):
    document.add_heading("4.6. Cơ chế Generate Prompt", level=1)
    p = document.add_paragraph()
    p.add_run(
        "Generate Prompt giúp người dùng chuyển từng finding thành một prompt sửa lỗi có thể đưa cho coding assistant. Cơ chế này nằm ở cấp từng issue, không phải ở cấp toàn bộ repository."
    )
    set_paragraph_font(p)

    add_numbered(
        document,
        [
            "Người dùng mở report và chọn một issue code smell hoặc security.",
            "FE gọi API generate prompt với roastId và issueId.",
            "Backend kiểm tra roastRun thuộc user hiện tại và report đã hoàn thành.",
            "Nếu issue chưa có suggestion, hệ thống sinh fallback suggestion theo type.",
            "Backend gọi Gemini để thiết kế fix prompt dựa trên repo, file, line range, severity, title, description và suggestion.",
            "Kết quả được lưu vào issue gồm fixPrompt, fixPromptVersion, fixPromptInstruction, fixPromptInstructionAnalysis và fixPromptGeneratedAt.",
            "FE hiển thị prompt cho người dùng copy hoặc tiếp tục regenerate.",
        ],
    )

    document.add_heading("Regenerate Prompt", level=2)
    p = document.add_paragraph()
    p.add_run(
        "Khi issue đã có prompt, hệ thống yêu cầu người dùng nhập regenerateInstruction. Instruction này được dùng làm yêu cầu ưu tiên cao nhất để tạo prompt mới, ví dụ muốn prompt ngắn hơn, chi tiết hơn, có diff cụ thể, tập trung vào test hoặc tập trung vào security impact. Backend yêu cầu AI phân tích instruction và trả về instructionAnalysis cùng fixPrompt mới."
    )
    set_paragraph_font(p)

    add_table(
        document,
        ["Trường dữ liệu", "Vai trò"],
        [
            ["fixPrompt", "Nội dung prompt cuối cùng để người dùng sử dụng."],
            ["fixPromptVersion", "Số phiên bản prompt, tăng sau mỗi lần generate/regenerate."],
            ["fixPromptInstruction", "Yêu cầu tùy chỉnh của người dùng khi regenerate."],
            ["fixPromptInstructionAnalysis", "Diễn giải của AI về yêu cầu regenerate."],
            ["fixPromptGeneratedAt", "Thời điểm prompt được tạo."],
        ],
        widths=[2.1, 4.3],
    )

    p = document.add_paragraph()
    p.add_run("Ý nghĩa nghiệp vụ. ").bold = True
    p.add_run(
        "Cơ chế này biến report từ danh sách lỗi tĩnh thành workflow hỗ trợ sửa lỗi. Người dùng không chỉ biết issue nằm ở đâu, mà còn có prompt bám sát file, line range, severity và hướng sửa cụ thể cho issue đó."
    )
    set_paragraph_font(p)


def build():
    document = Document()
    configure_document(document)
    add_title_page(document)
    add_section_42(document)
    add_section_43(document)
    add_section_44(document)
    add_section_45(document)
    add_section_46(document)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
