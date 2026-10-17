from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
DOCX_PATH = OUT_DIR / "vibecheck_repo_analysis_core_architecture.docx"
DIAGRAM_PATH = OUT_DIR / "vibecheck_repo_analysis_core_architecture.png"


COLORS = {
    "ink": "1F2937",
    "muted": "6B7280",
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "light_blue": "EAF3FB",
    "light_green": "EAF7EF",
    "light_yellow": "FFF7DB",
    "light_red": "FDEDEE",
    "light_gray": "F4F6F8",
    "border": "D0D7DE",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D0D7DE", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9360):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_dxa))


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width_dxa))


def set_cell_margins(cell, top=100, bottom=100, start=140, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, header_fill="F2F4F7", widths=None):
    set_table_width(table)
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[col_index])
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string(COLORS["ink"])


def add_paragraph(doc, text="", style=None, bold_prefix=None):
    paragraph = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def add_placeholder(doc, title, height_rows=8):
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COLORS["dark_blue"])

    table = doc.add_table(rows=height_rows, cols=1)
    set_table_width(table)
    for row in table.rows:
        row.height = Inches(0.22)
        cell = row.cells[0]
        set_cell_border(cell, color="B7C4D1")
        set_cell_shading(cell, "FBFCFE")
        set_cell_margins(cell, top=80, bottom=80, start=160, end=160)
    cell = table.cell(0, 0)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Paste screenshot here")
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(COLORS["muted"])
    doc.add_paragraph()


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_border(cell, color="9CC7E8")
    set_cell_shading(cell, "EAF3FB")
    set_cell_margins(cell, top=130, bottom=130, start=180, end=180)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(label)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COLORS["dark_blue"])
    paragraph.add_run(f" {text}")
    doc.add_paragraph()


def create_diagram():
    image = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(image)

    def font(size, bold=False):
        candidates = [
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
            "/System/Library/Fonts/Supplemental/Helvetica Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Helvetica.ttf",
        ]
        for candidate in candidates:
            try:
                return ImageFont.truetype(candidate, size)
            except Exception:
                continue
        return ImageFont.load_default()

    title_font = font(34, bold=True)
    box_font = font(21)
    small_font = font(18)
    ink = "#1F2937"
    border = "#8EA2B7"
    arrow = "#64748B"

    draw.text(
        (900, 48),
        "VibeCheck Core Repository Analysis Architecture",
        anchor="mm",
        font=title_font,
        fill="#0F172A",
    )

    def draw_centered_text(rect, text, fill=ink):
        x1, y1, x2, y2 = rect
        lines = text.split("\n")
        line_heights = []
        line_widths = []
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=box_font)
            line_widths.append(bbox[2] - bbox[0])
            line_heights.append(bbox[3] - bbox[1])
        total_height = sum(line_heights) + (len(lines) - 1) * 8
        y = y1 + ((y2 - y1) - total_height) / 2
        for line, width, height in zip(lines, line_widths, line_heights):
            draw.text(
                (x1 + ((x2 - x1) - width) / 2, y),
                line,
                font=box_font if len(lines) <= 2 else small_font,
                fill=fill,
            )
            y += height + 8

    def box(rect, text, fill):
        draw.rounded_rectangle(rect, radius=18, fill=f"#{fill}", outline=border, width=3)
        draw_centered_text(rect, text)

    def arrow_line(start, end):
        draw.line([start, end], fill=arrow, width=4)
        sx, sy = start
        ex, ey = end
        dx = ex - sx
        dy = ey - sy
        length = max((dx * dx + dy * dy) ** 0.5, 1)
        ux = dx / length
        uy = dy / length
        left = (-uy, ux)
        size = 16
        base = (ex - ux * size, ey - uy * size)
        points = [
            (ex, ey),
            (base[0] + left[0] * size * 0.55, base[1] + left[1] * size * 0.55),
            (base[0] - left[0] * size * 0.55, base[1] - left[1] * size * 0.55),
        ]
        draw.polygon(points, fill=arrow)

    boxes = {
        "frontend": ((70, 160, 345, 260), "Frontend\nReact UI", COLORS["light_blue"]),
        "api": ((425, 160, 720, 260), "Auth + REST API\nExpress", COLORS["light_blue"]),
        "repo": ((800, 160, 1095, 260), "Repo Intake\nGitHub API", COLORS["light_green"]),
        "files": ((1180, 160, 1535, 260), "File Selector\nIgnore + priority", COLORS["light_green"]),
        "cache": ((90, 430, 430, 540), "Cache Lookup\nfile SHA + rules hash", COLORS["light_yellow"]),
        "agents": ((510, 410, 1265, 565), "Domain Multi-Agent Analysis\nMaintainability | Architecture | Security\nRobustness | Operational | Interaction", "F3E8FF"),
        "ai": ((1355, 430, 1665, 540), "AI Provider\nGemini / OpenAI", COLORS["light_red"]),
        "validation": ((425, 680, 765, 790), "Validation Phase\nprecision filter", COLORS["light_yellow"]),
        "score": ((875, 680, 1185, 790), "Score + Grade\nP0-P3 summary", COLORS["light_blue"]),
        "mongo": ((1295, 680, 1635, 790), "MongoDB\nruns, issues, cache", COLORS["light_green"]),
        "response": ((690, 865, 1110, 950), "Report API Response\nhistory + detail + prompts", COLORS["light_blue"]),
    }
    for rect, text, fill in boxes.values():
        box(rect, text, fill)

    arrows = [
        ((345, 210), (425, 210)),
        ((720, 210), (800, 210)),
        ((1095, 210), (1180, 210)),
        ((1340, 260), (1040, 410)),
        ((865, 260), (265, 430)),
        ((430, 485), (510, 485)),
        ((1265, 485), (1355, 485)),
        ((880, 565), (595, 680)),
        ((765, 735), (875, 735)),
        ((1185, 735), (1295, 735)),
        ((1465, 680), (1000, 865)),
        ((1030, 790), (950, 865)),
    ]
    for start, end in arrows:
        arrow_line(start, end)

    draw.text(
        (85, 942),
        "Simplified view: exact implementation runs inside one Express backend with shared concurrency, cache, and provider configuration.",
        font=font(17),
        fill="#64748B",
    )
    image.save(DIAGRAM_PATH)


def setup_document_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, COLORS["blue"], 16, 8),
        ("Heading 2", 13, COLORS["blue"], 12, 6),
        ("Heading 3", 12, COLORS["dark_blue"], 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build_document():
    create_diagram()

    doc = Document()
    setup_document_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("VibeCheck Repository Analysis Core Architecture")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run(
        "Technical workflow summary for the current backend analysis pipeline"
    )
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(COLORS["muted"])

    add_callout(
        doc,
        "Purpose.",
        "This document explains the core backend architecture used to check a GitHub repository, run AI-assisted multi-domain analysis, persist findings, and return report data to the frontend.",
    )

    doc.add_heading("1. Scope", level=1)
    add_paragraph(
        doc,
        "The current VibeCheck backend is implemented with Node.js, Express.js, MongoDB, and Mongoose. It authenticates users through GitHub OAuth and analyzes public GitHub repositories through a domain-based multi-agent pipeline. The goal of the core pipeline is to produce a trustworthy repository review, not merely a fast response.",
    )
    add_bullet(
        doc,
        "Repository checking uses GitHub metadata, the default branch, file tree, file count, size, and deterministic repository size classification.",
    )
    add_bullet(
        doc,
        "Repository analysis uses six domain agents: Maintainability, Architecture, Security, Robustness, Operational, and Interaction.",
    )
    add_bullet(
        doc,
        "The backend stores completed runs, issue rows, per-file AI cache entries, and daily analysis usage events in MongoDB.",
    )

    doc.add_heading("2. Core Architecture", level=1)
    doc.add_picture(str(DIAGRAM_PATH), width=Inches(6.5))
    caption = doc.add_paragraph("Figure 1. Simplified backend architecture and repository analysis flow.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.italic = True
        run.font.color.rgb = RGBColor.from_string(COLORS["muted"])

    doc.add_heading("3. Main Backend Components", level=1)
    table = doc.add_table(rows=1, cols=3)
    widths = [2200, 3000, 4160]
    headers = ["Component", "Main responsibility", "Output"]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    rows = [
        (
            "Repository Check API",
            "Validates the GitHub URL, fetches repository metadata, reads the tree, applies ignore rules, and classifies size.",
            "roastId, repository snapshot, file count, size label, next action.",
        ),
        (
            "Repository Ingestion",
            "Loads prioritized source files from the default branch while ignoring generated, dependency, binary, media, build, and hidden noise paths.",
            "Filtered code sample with file path, size, SHA, priority score, and source content.",
        ),
        (
            "Domain Agents",
            "Runs six specialized analysis domains using detection and validation prompts.",
            "Validated issue lists grouped by Maintainability, Architecture, Security, Robustness, Operational, and Interaction.",
        ),
        (
            "AI Provider Layer",
            "Calls Gemini for system analysis or OpenAI when a user supplies a personal API key.",
            "Strict JSON result parsed into normalized findings.",
        ),
        (
            "Cache Layer",
            "Stores raw per-file detections by repository name, file path, GitHub file SHA, provider, model, and rules hash.",
            "Reusable raw detections for unchanged files.",
        ),
        (
            "Persistence Layer",
            "Stores runs, issues, repository metadata, cache entries, and daily usage records in MongoDB.",
            "History, detail reports, issue prompts, and quota enforcement.",
        ),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, widths=widths)

    doc.add_heading("4. Repository Analysis Flow", level=1)
    steps = [
        "The user signs in with GitHub OAuth and receives a JWT for authenticated API calls.",
        "The frontend submits a public GitHub repository URL to POST /api/repos/check.",
        "The backend validates the URL, fetches GitHub metadata, reads the default branch tree, applies ignore rules, and creates a checked roast run.",
        "The frontend calls POST /api/roasts/{roastId}/analyze. If the user supplies an OpenAI key, the provider becomes user_openai; otherwise the system Gemini provider is used.",
        "The backend loads prioritized source files. Dependency folders, build outputs, lockfiles, minified files, binary media, and irrelevant hidden folders are excluded.",
        "Each domain agent performs a high-recall detection phase and a precision-focused validation phase.",
        "The backend normalizes findings into a common issue format, assigns P0-P3 priority, computes score and grade, stores issues, and returns the report.",
        "The frontend renders the report, grouped issues, issue detail, GitHub file links, suggestions, and generated repair prompts.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("4.1 Core Analysis Guarantees", level=2)
    add_bullet(
        doc,
        "Completed reports are expected to come from AI-backed detection and validation, not silent local fallback, when strict mode is enabled.",
    )
    add_bullet(
        doc,
        "Every final issue includes a domain type, file path, GitHub file URL, line range, P0-P3 priority, explanation, and fix suggestion.",
    )
    add_bullet(
        doc,
        "Unchanged files can reuse cached raw detections when the repository blob SHA, model, provider, and rules hash match.",
    )
    add_bullet(
        doc,
        "If the provider returns quota, timeout, unavailable model, or invalid JSON errors after retries, the backend returns a structured failure instead of saving an incomplete report.",
    )

    doc.add_page_break()
    doc.add_heading("5. Domain-Based Multi-Agent Analysis", level=1)
    table = doc.add_table(rows=1, cols=3)
    for idx, header in enumerate(["Agent", "Focus", "Typical findings"]):
        table.rows[0].cells[idx].text = header
    domain_rows = [
        (
            "Maintainability",
            "Code smell and long-term changeability.",
            "Large functions, mixed responsibilities, duplicated logic, unclear naming, oversized components.",
        ),
        (
            "Architecture",
            "System structure and module boundaries.",
            "Tight coupling, controller-service mixing, weak layering, circular dependency risk.",
        ),
        (
            "Security",
            "Production security and trust boundaries.",
            "Injection risk, unsafe auth checks, exposed secrets, XSS patterns, unsafe external input handling.",
        ),
        (
            "Robustness",
            "Failure handling and edge-case resilience.",
            "Missing validation, null/undefined hazards, weak async error handling, brittle state assumptions.",
        ),
        (
            "Operational",
            "Runtime readiness and deployment safety.",
            "Missing health checks, risky environment config, weak logging, missing quotas, fragile deployment assumptions.",
        ),
        (
            "Interaction",
            "User-facing and API interaction quality.",
            "Poor loading/error states, weak form validation, confusing API response behavior, inconsistent interaction paths.",
        ),
    ]
    for row in domain_rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, header_fill="E8EEF5", widths=[1900, 3000, 4460])

    doc.add_heading("6. Detection and Validation Strategy", level=1)
    add_paragraph(
        doc,
        "Each domain follows a two-phase workflow. The detection phase is intentionally inclusive and aims for high recall. It may include uncertain findings so that suspicious code is not missed early. The validation phase is stricter and removes weak, duplicated, or unsupported findings before the final report is saved.",
    )
    add_bullet(
        doc,
        "Detection phase: analyzes selected file batches with domain-specific Rule 1 prompts and returns raw detections.",
    )
    add_bullet(
        doc,
        "Validation phase: re-checks all raw detections with domain-specific Rule 2 prompts and returns clean final JSON.",
    )
    add_bullet(
        doc,
        "Strict mode: local heuristic fallback is disabled when report accuracy matters. If the AI provider fails, the backend returns an error instead of saving an incomplete fallback report.",
    )

    doc.add_heading("7. Scoring and Priority Classification", level=1)
    table = doc.add_table(rows=1, cols=3)
    for idx, header in enumerate(["Priority", "Meaning", "Expected action"]):
        table.rows[0].cells[idx].text = header
    priority_rows = [
        ("P0", "Production blocker", "Must be fixed before deployment."),
        ("P1", "Must address before scaling", "May not block the MVP, but should be fixed before significant growth."),
        ("P2", "Should address during continued development", "Add to the near-term development backlog."),
        ("P3", "Acceptable MVP debt", "Accept temporarily, monitor, and revisit later."),
    ]
    for row in priority_rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, header_fill="F2F4F7", widths=[1300, 3100, 4960])
    add_paragraph(
        doc,
        "The final score starts from 100 and applies weighted penalties based on issue priority and the domain penalty multiplier. The numeric score is converted into a grade from A to F for simple frontend display.",
    )

    doc.add_page_break()
    doc.add_heading("8. Persistence Model", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "MongoDB collection"
    table.rows[0].cells[1].text = "Stored information"
    persistence_rows = [
        ("repositories", "Repository owner, name, full name, GitHub URL, default branch, and last known size metadata."),
        ("roast_runs", "One repository check/analyze run, including status, check metadata, AI provider, score, grade, and metrics."),
        ("roast_issues", "Final validated issues, domain type, file path, line range, P0-P3 priority, suggestion, and generated fix prompt fields."),
        ("analysis_file_caches", "Raw per-file detections keyed by repository, file path, GitHub blob SHA, provider, model, and rules hash."),
        ("analysis_usages", "Daily analyze usage events used to enforce the per-user analysis limit."),
    ]
    for row in persistence_rows:
        cells = table.add_row().cells
        cells[0].text = row[0]
        cells[1].text = row[1]
    style_table(table, header_fill="E8EEF5", widths=[2400, 6960])

    doc.add_heading("9. API Surface Used by the Frontend", level=1)
    add_bullet(doc, "POST /api/repos/check: checks repository metadata and creates a checked roast run.")
    add_bullet(doc, "POST /api/roasts/{roastId}/analyze: runs AI analysis and returns the completed report.")
    add_bullet(doc, "GET /api/roasts/history: lists completed analysis history for the authenticated user.")
    add_bullet(doc, "GET /api/roasts/{roastId}: returns a saved report detail.")
    add_bullet(doc, "POST /api/roasts/{roastId}/issues/{issueId}/prompt: generates or regenerates a fix prompt for one issue.")
    add_bullet(doc, "DELETE /api/roasts/{roastId}: removes one saved report from the user's history.")

    doc.add_page_break()
    doc.add_heading("10. Current Quality-Oriented Runtime Settings", level=1)
    add_paragraph(
        doc,
        "The current local configuration favors correctness over speed. AI fallback is disabled, retry-after responses from Gemini are respected, and request concurrency is reduced to avoid provider rate-limit spikes.",
    )
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Environment variable"
    table.rows[0].cells[1].text = "Purpose"
    settings_rows = [
        ("ANALYSIS_ALLOW_LOCAL_FALLBACK=false", "Prevents incomplete local fallback reports from being saved as completed AI reports."),
        ("ANALYSIS_MAX_CONCURRENT=2", "Limits total concurrent AI requests."),
        ("ANALYSIS_DETECTION_BATCH_SIZE=2", "Keeps prompts smaller and easier for the model to return as valid JSON."),
        ("ANALYSIS_DOMAIN_MAX_FILES=150", "Allows broader repository coverage per domain."),
        ("LLM_RETRY_COUNT=8", "Retries transient provider failures and invalid JSON responses."),
        ("LLM_REQUEST_TIMEOUT_MS=120000", "Allows longer model processing time per request."),
        ("LLM_MAX_RETRY_AFTER_MS=120000", "Caps provider retry-after waits to two minutes per retry."),
    ]
    for row in settings_rows:
        cells = table.add_row().cells
        cells[0].text = row[0]
        cells[1].text = row[1]
    style_table(table, widths=[3300, 6060])

    doc.add_page_break()
    doc.add_heading("11. Screenshot Evidence Placeholders", level=1)
    add_paragraph(
        doc,
        "Use this section to paste screenshots from the innspill_frontend repository review after the backend and frontend flows are tested.",
    )
    add_placeholder(
        doc,
        "Screenshot Placeholder 1: Backend API result for innspill_frontend",
        height_rows=9,
    )
    doc.add_page_break()
    add_placeholder(
        doc,
        "Screenshot Placeholder 2: Frontend report screen for innspill_frontend",
        height_rows=9,
    )
    doc.add_page_break()
    add_placeholder(
        doc,
        "Screenshot Placeholder 3: Issue detail view after selecting one issue",
        height_rows=9,
    )

    doc.add_page_break()
    doc.add_heading("12. Summary", level=1)
    add_paragraph(
        doc,
        "VibeCheck's core repository analysis architecture is a backend-centered, domain-based multi-agent pipeline. It ingests a public GitHub repository, filters and prioritizes analyzable files, runs specialized AI analysis domains, validates findings, scores the repository, persists the report, and exposes the result through REST APIs consumed by the frontend.",
    )

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build_document())
