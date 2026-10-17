from pathlib import Path
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "VibeCheck_AI_Analysis_Methodology_and_Prompt_Design.docx"
DIAGRAM = Path(tempfile.gettempdir()) / "vibecheck_ai_pipeline_diagram.png"


BLACK = "000000"
DARK_BLUE = "1F4E78"
MID_BLUE = "5B9BD5"
LIGHT_BLUE = "DDEBF7"
PALE_BLUE = "EEF5FB"
LIGHT_GRAY = "F2F2F2"
BORDER_GRAY = "D9D9D9"
WHITE = "FFFFFF"
PALE_GREEN = "E2F0D9"
PALE_YELLOW = "FFF2CC"
PALE_PINK = "FCE4D6"
PALE_PURPLE = "E4DFEC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def set_run_font(run, name="Aptos", size=None, bold=None, color=BLACK, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    doc.styles["Caption"].font.name = "Aptos"
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = True
    doc.styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    add_page_field(p)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def add_paragraph(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"{index}.  {item}")
        set_run_font(run)


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True)
    return p


def add_code_block(doc, text, font_size=8.3):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    p_pr.append(keep_lines)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F7F7F7")
    p_pr.append(shd)
    for index, line in enumerate(text.strip("\n").splitlines()):
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=font_size, color="202020")
        if index < len(text.strip("\n").splitlines()) - 1:
            run.add_break()
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for col_index, value in enumerate(headers):
        cell = header.cells[col_index]
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value)
        set_run_font(run, size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[col_index])

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            set_cell_shading(cell, PALE_BLUE if row_index % 2 else WHITE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
            if widths:
                cell.width = Inches(widths[col_index])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def rounded_box(draw, xy, fill, outline, title, subtitle, font_title, font_body):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=3)
    title_bbox = draw.textbbox((0, 0), title, font=font_title)
    tx = x1 + (x2 - x1 - (title_bbox[2] - title_bbox[0])) / 2
    draw.text((tx, y1 + 24), title, fill="#1F1F1F", font=font_title)
    lines = subtitle.split("\n")
    y = y1 + 72
    for line in lines:
        box = draw.textbbox((0, 0), line, font=font_body)
        lx = x1 + (x2 - x1 - (box[2] - box[0])) / 2
        draw.text((lx, y), line, fill="#3F3F3F", font=font_body)
        y += 30


def arrow(draw, start, end, color="#6B7280", width=5):
    draw.line([start, end], fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) > abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 18 * direction, ey - 10), (ex - 18 * direction, ey + 10)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 10, ey - 18 * direction), (ex + 10, ey - 18 * direction)]
    draw.polygon(points, fill=color)


def make_diagram(path):
    image = Image.new("RGB", (1800, 1130), "white")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 31)
        body_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 25)
        heading_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 42)
    except OSError:
        title_font = ImageFont.load_default()
        body_font = ImageFont.load_default()
        heading_font = ImageFont.load_default()

    heading = "VibeCheck repository analysis pipeline"
    heading_box = draw.textbbox((0, 0), heading, font=heading_font)
    draw.text(((1800 - (heading_box[2] - heading_box[0])) / 2, 35), heading, fill="#000000", font=heading_font)

    boxes = [
        ((70, 145, 390, 305), "#DDEBF7", "#5B9BD5", "Repository check", "GitHub metadata\nrecursive tree\nsize classification"),
        ((510, 145, 830, 305), "#E2F0D9", "#70AD47", "File preparation", "ignore rules\npriority scoring\ncontent download"),
        ((950, 145, 1270, 305), "#FFF2CC", "#BF9000", "Cache lookup", "blob SHA + domain\nprovider + model\nrules hash"),
        ((1390, 145, 1710, 305), "#FCE4D6", "#C55A11", "Detection batches", "Rule 1\nhigh recall\nstructured JSON"),
        ((1390, 445, 1710, 605), "#E4DFEC", "#8064A2", "Validation", "Rule 2\nhigh precision\ndeduplicate"),
        ((950, 445, 1270, 605), "#DDEBF7", "#5B9BD5", "Normalize findings", "P0 to P3\nfile and line links\nsuggestions"),
        ((510, 445, 830, 605), "#E2F0D9", "#70AD47", "Merge and score", "six domains\nweighted penalties\ngrade A to F"),
        ((70, 445, 390, 605), "#FFF2CC", "#BF9000", "Persist and return", "roast run\nissue records\nhistory and detail"),
    ]
    for xy, fill, outline, title, subtitle in boxes:
        rounded_box(draw, xy, fill, outline, title, subtitle, title_font, body_font)

    for start, end in [
        ((390, 225), (510, 225)),
        ((830, 225), (950, 225)),
        ((1270, 225), (1390, 225)),
        ((1550, 305), (1550, 445)),
        ((1390, 525), (1270, 525)),
        ((950, 525), (830, 525)),
        ((510, 525), (390, 525)),
    ]:
        arrow(draw, start, end)

    domain_y = 735
    draw.text((70, 670), "Six specialized analysis tasks share one concurrency limiter", fill="#000000", font=title_font)
    domain_labels = [
        ("Maintainability", "#DDEBF7", "#5B9BD5"),
        ("Architecture", "#E2F0D9", "#70AD47"),
        ("Security", "#FCE4D6", "#C55A11"),
        ("Robustness", "#FFF2CC", "#BF9000"),
        ("Operational", "#E4DFEC", "#8064A2"),
        ("Interaction", "#DDEBF7", "#5B9BD5"),
    ]
    x = 70
    for label, fill, outline in domain_labels:
        draw.rounded_rectangle((x, domain_y, x + 255, domain_y + 95), radius=20, fill=fill, outline=outline, width=3)
        box = draw.textbbox((0, 0), label, font=body_font)
        draw.text((x + (255 - (box[2] - box[0])) / 2, domain_y + 31), label, fill="#1F1F1F", font=body_font)
        x += 285

    draw.text((70, 885), "Provider path", fill="#000000", font=title_font)
    rounded_box(draw, (70, 930, 565, 1080), "#E2F0D9", "#70AD47", "System Gemini", "Default for small and medium\noptional local fallback", title_font, body_font)
    rounded_box(draw, (650, 930, 1145, 1080), "#DDEBF7", "#5B9BD5", "Personal OpenAI key", "Selected when a key is supplied\nrequired for large by default", title_font, body_font)
    rounded_box(draw, (1230, 930, 1730, 1080), "#F2F2F2", "#7F7F7F", "Shared controls", "JSON output, timeout, retry\nand bounded concurrency", title_font, body_font)
    image.save(path, quality=95)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(56)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(16)
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)
    run = title.add_run("VibeCheck AI Analysis Methodology and Prompt Design")
    set_run_font(run, name="Aptos Display", size=27, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("Technical system report for reproducibility and reviewer assessment")
    set_run_font(run, size=14, color="404040")

    add_paragraph(
        doc,
        "This report explains how VibeCheck ingests a public GitHub repository, selects source files, constructs AI prompts, performs two-phase issue detection and validation across six quality domains, classifies production priority, stores results, and generates repair prompts. It includes implementation-aligned prompt excerpts and rule structures so that readers can assess how the language models are instructed to identify production debt.",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    run = p.add_run("System scope")
    set_run_font(run, bold=True, size=11)
    add_bullets(
        doc,
        [
            "Backend: Node.js with Express and MongoDB through Mongoose.",
            "Repository source: public GitHub repositories accessed through the GitHub API and raw content endpoints.",
            "AI providers: a system Gemini model or a user-supplied OpenAI API key.",
            "Analysis domains: Maintainability, Architecture, Security, Robustness, Operational readiness, and Interaction reliability.",
            "Output: structured issue records, P0-P3 production priority, repository score and grade, history, issue-specific suggestions, and reusable repair prompts.",
        ],
    )
    doc.add_page_break()


def build_document():
    make_diagram(DIAGRAM)
    doc = Document()
    style_document(doc)
    add_cover(doc)

    add_heading(doc, "Executive Summary", 1)
    add_paragraph(
        doc,
        "VibeCheck uses a hybrid pipeline rather than asking one language model to judge an entire repository in a single request. Deterministic code first validates the repository, builds a recursive file inventory, applies ignore rules, classifies repository size, prioritizes analyzable files, and downloads bounded source snippets. The AI layer then runs six specialized analysis tasks. Each task performs a high-recall Detection Phase under Rule 1, followed by a high-precision Validation Phase under Rule 2. The final findings are normalized, assigned P0-P3 production priorities, merged, scored, stored in MongoDB, and returned through the backend API.",
    )
    add_paragraph(
        doc,
        "The implementation is best described as a lightweight multi-agent analysis pattern. The six domain analyzers have different roles, rule documents, focus areas, heuristic metrics, and output contracts, but they are not autonomous agents with independent tools or memory. They share the same provider client, bounded concurrency limiter, cache, normalizer, and report merger. This distinction matters when evaluating speed, reproducibility, and the scope of the claimed contribution.",
    )
    add_paragraph(
        doc,
        "The prompts are deliberately split into recall and precision objectives. Phase 1 is instructed to include plausible issues and attach evidence. Phase 2 receives only the raw detections and the stricter validation rules; it removes weak findings, merges duplicates, and emits a stable JSON schema. This design makes the central AI capability inspectable: the model is not merely told to find production debt, but is constrained by domain definitions, evidence requirements, metric hints, output schemas, and explicit filtering criteria.",
    )

    add_heading(doc, "Contents", 1)
    add_numbered(
        doc,
        [
            "System architecture and processing flow",
            "Repository ingestion and deterministic preprocessing",
            "Six-domain AI analysis design",
            "Two-phase prompt methodology",
            "Domain-specific rules and prompt construction",
            "AI provider integration and structured output",
            "Caching, concurrency, retry, and fallback behavior",
            "Priority classification, scoring, and report generation",
            "Repair prompt generation and regeneration",
            "Persistence model and API usage",
            "Reproducibility, limitations, and reviewer assessment",
            "Appendix with prompt skeletons and implementation map",
        ],
    )
    doc.add_page_break()

    add_heading(doc, "1 System Architecture and Processing Flow", 1)
    add_paragraph(
        doc,
        "The backend separates deterministic repository processing from probabilistic AI reasoning. GitHub access, file filtering, size classification, prioritization, caching, concurrency, normalization, scoring, persistence, and authorization are implemented in application code. The LLM is used only for domain-specific interpretation, validation, and repair-prompt synthesis.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(DIAGRAM), width=Inches(6.72))
    add_caption(doc, "Figure 1. Backend repository analysis architecture and AI processing flow")

    add_heading(doc, "1.1 Component responsibilities", 2)
    add_table(
        doc,
        ["Component", "Primary responsibility", "Deterministic or AI"],
        [
            ["Repository check", "Validate a public GitHub URL; fetch repository, branch, commit, and recursive tree metadata.", "Deterministic"],
            ["File preparation", "Apply ignore rules, rank files, enforce file-count and byte limits, and download source content.", "Deterministic"],
            ["Domain analyzers", "Detect and validate issues in six specialized quality domains.", "AI-assisted"],
            ["File cache", "Reuse raw Phase 1 detections for unchanged Git blobs under the same provider, model, domain, and rules.", "Deterministic"],
            ["Normalizer", "Map variable provider output into stable issue records and context-specific suggestions.", "Deterministic"],
            ["Scoring engine", "Apply P0-P3 weights and domain multipliers; calculate score and grade.", "Deterministic"],
            ["Prompt generator", "Create or regenerate an issue-specific repair prompt for use in a coding assistant.", "AI-assisted with local fallback for first generation"],
            ["Persistence and API", "Store repository runs, issues, usage quota, and cache; expose analysis, history, detail, delete, and prompt endpoints.", "Deterministic"],
        ],
        widths=[1.35, 4.15, 1.2],
    )

    add_heading(doc, "1.2 End-to-end sequence", 2)
    add_numbered(
        doc,
        [
            "The authenticated user submits a public GitHub repository URL to POST /api/repos/check.",
            "The backend validates the URL, fetches GitHub metadata and the recursive tree, applies ignore rules, calculates deterministic size statistics, and creates a checked roast run.",
            "The user calls POST /api/roasts/{roastId}/analyze, optionally supplying a personal OpenAI API key.",
            "The backend enforces the daily analysis quota and provider policy, reloads the repository at the current default-branch commit, selects and downloads analyzable files, and creates a shared AI concurrency limiter.",
            "Six domain analyses start concurrently. Each domain selects relevant files, retrieves cached raw detections, analyzes cache misses in small batches, and validates the combined detections once.",
            "The backend normalizes and persists issue rows, calculates domain totals, P0-P3 totals, score, grade, file links, and report metadata.",
            "The frontend reads the report immediately or retrieves it later through history and report-detail endpoints.",
            "For any issue, the user can generate a repair prompt and later regenerate it by specifying how the prompt should change.",
        ],
    )

    add_heading(doc, "2 Repository Ingestion and Deterministic Preprocessing", 1)
    add_heading(doc, "2.1 GitHub metadata collection", 2)
    add_paragraph(
        doc,
        "The check step accepts only github.com repository URLs with an owner and repository name. The backend requests repository metadata, the default branch, and a recursive Git tree. Private or unavailable repositories are rejected. A backend GitHub token can be configured to avoid the low unauthenticated rate limit commonly encountered on shared deployment infrastructure.",
    )

    add_heading(doc, "2.2 Ignore policy", 2)
    add_paragraph(
        doc,
        "The analyzer removes files that are unlikely to contain reviewable source code or that would consume tokens without improving findings. Generated output, dependency directories, caches, binary and media files, lock files, minified bundles, logs, and hidden directories are ignored. Hidden paths are excluded by default, with explicit exceptions for security- and deployment-relevant material.",
    )
    add_table(
        doc,
        ["Group", "Examples", "Treatment"],
        [
            ["Dependencies and build output", "node_modules, dist, build, .next, target, vendor, coverage", "Ignored"],
            ["Tool and editor state", ".git, .cache, .turbo, .idea, .vscode, __pycache__", "Ignored"],
            ["Generated and lock files", "package-lock.json, yarn.lock, pnpm-lock.yaml, Cargo.lock, *.min.js", "Ignored"],
            ["Binary and media", "images, archives, video, audio, fonts, executables, maps", "Ignored"],
            ["Security and deployment exceptions", ".github/workflows, Dockerfile, docker-compose.yml, .env.example", "Retained"],
            ["Reviewable source", "JS, TS, Python, Java, Go, C family, Vue, Svelte, HTML, CSS, SQL, YAML and related text files", "Eligible within size limits"],
        ],
        widths=[1.55, 3.3, 1.85],
    )

    add_heading(doc, "2.3 Deterministic repository size classification", 2)
    add_paragraph(
        doc,
        "Repository size is no longer delegated to an LLM. The backend calculates it from the filtered tree so repeated checks of the same commit produce the same class. Reviewable files and bytes are the primary signals; total retained files and bytes and GitHub tree truncation act as safeguards.",
    )
    add_table(
        doc,
        ["Class", "Conditions after ignore rules", "Provider policy"],
        [
            ["Small", "At most 50 reviewable files, 500,000 reviewable bytes, and 250 retained files.", "System Gemini by default; personal OpenAI key is optional."],
            ["Medium", "Exceeds a small threshold but has at most 250 reviewable files, 2,500,000 reviewable bytes, 1,500 retained files, and 75,000,000 retained bytes.", "System Gemini by default; personal OpenAI key is optional."],
            ["Large", "Exceeds any medium threshold or the GitHub recursive tree is truncated.", "Personal OpenAI key required by default; a controlled environment override can allow system Gemini."],
        ],
        widths=[0.75, 3.95, 2.0],
    )

    add_heading(doc, "2.4 File prioritization and coverage", 2)
    add_paragraph(
        doc,
        "The loader follows a priority-first strategy. It ranks files by path, language, security relevance, and size, then downloads at most the configured global limit. High-priority locations include pages, components, services, API clients, hooks, controllers, routes, middleware, authentication code, utilities, workflows, and deployment files. Constants, types, styles, mocks, fixtures, and tests receive lower scores. Each domain then reranks the downloaded snippets using domain-specific path hints and selects up to its own limit.",
    )
    add_table(
        doc,
        ["Configuration", "Default", "Effect"],
        [
            ["ANALYSIS_MAX_PRIORITY_FILES", "120", "Maximum source files downloaded after global prioritization."],
            ["ANALYSIS_MAX_FILE_BYTES", "80,000", "Files larger than this are not sent for analysis."],
            ["ANALYSIS_FETCH_CONCURRENT", "5", "Maximum concurrent raw-file downloads."],
            ["ANALYSIS_DOMAIN_MAX_FILES", "50", "Maximum files selected for one quality domain."],
            ["ANALYSIS_DETECTION_BATCH_SIZE", "3", "Files grouped into one Phase 1 LLM request."],
            ["ANALYSIS_MAX_CONCURRENT", "5", "Shared maximum concurrent AI requests across all domains and phases."],
        ],
        widths=[2.1, 0.8, 3.8],
    )
    add_paragraph(
        doc,
        "Coverage statement: VibeCheck does not claim exhaustive formal verification. It analyzes a bounded, prioritized subset of eligible files. This makes whole-project review practical while preserving a clear limitation for very large repositories.",
        bold_lead="Coverage statement:",
    )

    add_heading(doc, "3 Six Domain AI Analysis Design", 1)
    add_paragraph(
        doc,
        "The orchestrator starts all six domain analyzers with Promise.all. A single shared concurrency limiter prevents the combined workload from exceeding the configured request limit. Therefore, the domains can make progress in parallel, but the provider is protected from an uncontrolled burst of requests.",
    )
    add_table(
        doc,
        ["Domain", "Detection focus", "Heuristic signals", "Penalty multiplier"],
        [
            ["Maintainability", "Code smells, duplication, long methods, large components, dead code, naming, mixed responsibilities, coupling.", "Method length, class size, nesting, external dependency ratio, duplication.", "3"],
            ["Architecture", "Module boundaries, layering, dependency direction, cohesion, coupling, misplaced responsibilities, drift.", "Boundary clarity, coupling, cohesion, dependency direction, layering risk.", "5"],
            ["Security", "OWASP-inspired injection, XSS, auth and authorization, secrets, cryptography, SSRF, traversal, insecure configuration.", "Tainted input, auth boundary, validation, dangerous APIs, sensitive exposure.", "7"],
            ["Robustness", "Input validation, error handling, null safety, async failure, retries, fallbacks, edge cases.", "Validation strength, error coverage, null safety, async handling, fallback behavior.", "4"],
            ["Operational", "Environment safety, deployment, health, logging, rate limits, CORS, CI/CD, runtime and cost risk.", "Configuration, observability, deployment readiness, runtime safety, quota risk.", "4"],
            ["Interaction", "Loading and error states, form flows, API contract alignment, navigation, user feedback.", "Feedback quality, loading and error coverage, contract alignment, flow resilience.", "3"],
        ],
        widths=[1.0, 3.05, 2.15, 0.6],
        font_size=8.1,
    )

    add_heading(doc, "3.1 Lightweight multi-agent interpretation", 2)
    add_paragraph(
        doc,
        "Each domain is configured as a specialized role with its own Rule 1 file, Rule 2 file, issue noun, raw issue-name field, detection focus, validation focus, heuristic metrics, path hints, and score multiplier. The architecture provides specialization and concurrent execution without introducing autonomous planning, cross-agent conversation, or separate API credentials for each domain. One provider credential can serve every domain because requests are separated by prompt and rule context rather than by key.",
    )

    add_heading(doc, "4 Two Phase Prompt Methodology", 1)
    add_heading(doc, "4.1 Phase 1 Detection for high recall", 2)
    add_paragraph(
        doc,
        "Phase 1 is intentionally inclusive. The model is told to identify all plausible issues, include uncertain candidates, avoid aggressive filtering, attach concrete evidence, estimate heuristic metrics, and return raw JSON. The prompt is assembled from a common template plus domain-specific variables and the complete Rule 1 document. One prompt may contain a small batch of source files, each labeled with file name, relative path, byte size, and content.",
    )
    add_code_block(
        doc,
        """You are a <domain-specific senior role>.
Analyze the provided source code file(s) and detect ALL possible <domain issues>.

1. Strictly follow <DOMAIN_RULE_1.md>.
2. This is a DETECTION phase:
   - Be INCLUSIVE (high recall).
   - You may include uncertain issues.
   - Do not filter aggressively.
3. Focus on: <domain-specific detection focus>.
4. For every issue, provide rule-based reasoning, evidence, line range,
   an issue-specific suggestion, P0-P3 priority, and heuristic metrics.
5. Return only a RAW detection JSON list. Return [] when no issue is found.

<Complete Rule 1 content>
<File name, path, size, and source for each file in the batch>""",
    )

    add_heading(doc, "4.2 Raw detection contract", 2)
    add_code_block(
        doc,
        """[
  {
    "fileName": "<file name>",
    "filePath": "<relative path>",
    "<domainIssueNameField>": "<issue name>",
    "category": "<rule category>",
    "title": "<short issue title>",
    "description": "<why the issue is detected>",
    "suggestion": "<specific remediation for this issue and file>",
    "evidence": "<code snippet or structural evidence>",
    "startLine": 1,
    "endLine": 1,
    "priority": "P0 | P1 | P2 | P3",
    "metrics": { "<domain metrics>": "<estimated values>" }
  }
]""",
    )
    add_paragraph(
        doc,
        "Detections are matched back to the files in the submitted batch. Findings that reference an unknown path or unmatched file name are discarded before caching, reducing the risk that model output is attached to the wrong file.",
    )

    add_heading(doc, "4.3 Phase 2 Validation for high precision", 2)
    add_paragraph(
        doc,
        "Phase 2 receives the combined raw detections for one domain, including detections loaded from cache and detections produced for changed files. It does not re-read the full source batch. Instead, it re-evaluates evidence and reasoning under the stricter Rule 2 definitions. The model is told to remove weak, speculative, incorrect, or duplicate findings, merge findings that describe the same weakness, treat low-confidence findings as removable, improve the explanation and remediation, and return the final schema.",
    )
    add_code_block(
        doc,
        """You are a <domain-specific validation expert>.
RE-EVALUATE the RAW detections and strictly follow <DOMAIN_RULE_2.md>.

1. Focus on PRECISION.
2. Remove weak, incorrect, duplicate, or speculative issues.
3. Re-check the rule definition, evidence, reasoning, and metrics.
4. Internally classify confidence:
   high = strong evidence and clear rule match;
   medium = partial or borderline match;
   low = remove.
5. Preserve or improve the issue-specific suggestion.
6. Return clean, consistent, deduplicated FINAL JSON.
7. Use only P0, P1, P2, or P3. Return [] when no valid issue remains.

<Complete Rule 2 content>
<RAW detection JSON from Phase 1>""",
    )

    add_heading(doc, "4.4 Final validation contract", 2)
    add_code_block(
        doc,
        """[
  {
    "index": 1,
    "fileName": "<file name>",
    "filePath": "<relative path>",
    "title": "<validated issue title>",
    "description": "<precise validated explanation>",
    "suggestion": "<specific repair direction>",
    "startLine": 1,
    "endLine": 1,
    "priority": "P0 | P1 | P2 | P3",
    "priorityDisplay": "<priority and production meaning>"
  }
]""",
    )

    add_heading(doc, "4.5 Why the two phases are separated", 2)
    add_table(
        doc,
        ["Concern", "Detection Phase", "Validation Phase"],
        [
            ["Primary objective", "Recall: avoid missing plausible issues.", "Precision: remove weak or duplicate issues."],
            ["Allowed uncertainty", "Explicitly allowed.", "Low confidence is removed."],
            ["Evidence", "Collected from source and attached to each raw item.", "Re-checked for rule match and production relevance."],
            ["Metrics", "Estimated to support screening.", "Re-evaluated when needed."],
            ["Output", "Rich raw detections for internal processing and cache.", "Stable issue list for persistence and API response."],
        ],
        widths=[1.3, 2.7, 2.7],
    )

    add_heading(doc, "5 Domain Specific Rules and Prompt Construction", 1)
    add_paragraph(
        doc,
        "The common prompt skeleton controls reasoning behavior and output format. Domain Rule 1 and Rule 2 documents supply the substantive definitions. This avoids relying on a single vague request such as 'find production debt.' The following subsections summarize the instruction content that changes by domain.",
    )

    domain_sections = [
        (
            "5.1 Maintainability",
            "The Maintainability analyzer uses the former code-smell rules. Rule 1 defines Bloaters, Change Preventers, Couplers, Dispensables, and Object-Oriented Abusers. It includes thresholds such as methods longer than 50 lines, more than 10 methods for a large class, more than five parameters, and deep nesting. The detection prompt asks for method length, class size, nesting depth, external dependency ratio, and duplication signals. Rule 2 requires a formal rule match, strong observable evidence, supporting metrics, and consistent reasoning. It removes vague detections and merges duplicate smells in the same region.",
        ),
        (
            "5.2 Architecture",
            "The Architecture analyzer looks for boundary violations, layering problems, tight coupling, low cohesion, and architectural drift. The prompt role is a senior software architect specializing in modularity and boundary analysis. Rule 2 removes stylistic concerns and claims that require repository-wide knowledge not present in the supplied evidence. This prevents a file-level batch from making unsupported claims about the complete system architecture.",
        ),
        (
            "5.3 Security",
            "The Security analyzer uses an OWASP-inspired workflow but does not claim formal OWASP compliance or complete vulnerability verification. Rule 1 covers injection, cross-site scripting, authentication and authorization, secrets and sensitive data, cryptography, server-side request forgery, path traversal and file handling, and insecure configuration. The prompt asks the model to track untrusted input, trust boundaries, validation strength, dangerous API use, and sensitive-data exposure. Rule 2 retains a finding only when there is concrete evidence and a realistic exploit path or security weakness. Test-only and purely speculative items are removed, while public routes, request handlers, external calls, configuration, and secrets receive stricter review.",
        ),
        (
            "5.4 Robustness",
            "The Robustness analyzer targets invalid-input behavior, weak error handling, null and undefined hazards, missing retries or fallbacks, and unhandled edge cases. Rule 2 requires a concrete failure mode and removes generic requests for better error handling. Production-facing paths are preferred over test-only concerns.",
        ),
        (
            "5.5 Operational Readiness",
            "The Operational analyzer examines configuration safety, runtime safeguards, observability, deployment readiness, and cost or quota risk. Typical prompt targets include missing environment validation, localhost production paths, missing health checks, unsafe CORS, missing rate limits, incomplete CI or Docker configuration, and expensive AI calls without caching or concurrency control. Validation keeps only findings that can affect deployment, production runtime, monitoring, cost, or scaling.",
        ),
        (
            "5.6 Interaction Reliability",
            "The Interaction analyzer reviews user-visible and frontend-backend workflow behavior: loading and pending states, actionable error feedback, success feedback, API contract alignment, form validation, and callback or navigation failure handling. Rule 2 excludes purely aesthetic preferences unless they affect task completion and requires evidence of a real user action or API flow.",
        ),
    ]
    for heading, body in domain_sections:
        add_heading(doc, heading, 2)
        add_paragraph(doc, body)

    add_heading(doc, "5.7 Rule and prompt composition", 2)
    add_paragraph(
        doc,
        "For each request, the backend concatenates the behavioral template, domain metadata, full rule content, and file or raw-detection input. Rule files are therefore executable prompt context rather than documentation that is separate from the model call. Updating a Rule 1 document changes the detection cache key through a SHA-256 rule hash, so old raw detections are not reused under changed detection rules.",
    )

    add_heading(doc, "6 AI Provider Integration and Structured Output", 1)
    add_heading(doc, "6.1 Provider selection", 2)
    add_table(
        doc,
        ["Condition", "Provider", "Notes"],
        [
            ["No personal key and repository is small or medium", "system_gemini", "Uses the backend Gemini key."],
            ["Personal OpenAI key supplied for any repository size", "user_openai", "The user key is passed to the provider call; only a boolean indicating personal-key use is stored."],
            ["Large repository without personal key", "Rejected by default", "The controlled ALLOW_SYSTEM_AI_FOR_LARGE override can permit Gemini in a trusted environment."],
        ],
        widths=[2.6, 1.25, 2.85],
    )
    add_paragraph(
        doc,
        "The exact model is environment-configurable through GEMINI_ANALYSIS_MODEL and OPENAI_ANALYSIS_MODEL. The code provides fallback defaults, while deployment configuration can select another compatible model without changing the prompt pipeline.",
    )

    add_heading(doc, "6.2 Provider request controls", 2)
    add_bullets(
        doc,
        [
            "Gemini receives a user message containing the assembled prompt and requests JSON output.",
            "OpenAI receives a system instruction to act as a senior static-analysis assistant and return only JSON matching the requested schema, followed by the assembled user prompt.",
            "Both paths use a low temperature of 0.15 to reduce output variation.",
            "Timeout, retry count, and retry delay are configurable. Retryable statuses include request timeout, conflict, rate limit, and common transient server errors.",
            "The response parser removes optional Markdown JSON fences and rejects invalid JSON with a structured backend error.",
        ],
    )
    add_paragraph(
        doc,
        "The provider credential is not included in the prompt. A personal OpenAI key is used for the outgoing Authorization header during the analysis request and is not written to the roast run or issue records.",
    )

    add_heading(doc, "7 Caching Concurrency Retry and Fallback", 1)
    add_heading(doc, "7.1 Per-file detection cache", 2)
    add_paragraph(
        doc,
        "VibeCheck caches raw Phase 1 detections, not source code and not the final validated report. The cache is deliberately scoped so that a result can be reused only when the repository, path, Git blob, domain, provider, model, and Rule 1 version are unchanged.",
    )
    add_code_block(
        doc,
        """Cache key = {
  repositoryFullName,
  filePath,
  fileSha,
  analysisType,
  provider,
  model,
  rulesHash
}""",
    )
    add_paragraph(
        doc,
        "A changed file receives a new Git blob SHA and is analyzed again. Unchanged files can reuse raw detections. The validation phase still runs once on the full current set of raw detections so cached and new items are filtered together under the current report context.",
    )

    add_heading(doc, "7.2 Batching and bounded parallelism", 2)
    add_paragraph(
        doc,
        "Cache misses are grouped into small detection batches. Every batch and validation request passes through the same in-process concurrency limiter. Six domains are launched concurrently, but only the configured number of AI requests can be active. This reduces total wall-clock time without issuing an unbounded number of provider calls.",
    )

    add_heading(doc, "7.3 Failure behavior", 2)
    add_paragraph(
        doc,
        "Transient provider calls are retried according to configuration. When system Gemini detection fails with a provider, quota, timeout, server, or invalid-JSON error, the affected batch may use local heuristic fallback checks. Fallback detections are not cached. If Gemini validation fails, normalized raw detections can be returned under a fallback model marker. Personal OpenAI errors are surfaced to the caller rather than silently replaced, because the user explicitly selected that provider. The local fallback is a continuity mechanism and should not be treated as equivalent to a successful LLM review.",
    )

    add_heading(doc, "8 Priority Classification Scoring and Report Generation", 1)
    add_heading(doc, "8.1 Production priority", 2)
    add_table(
        doc,
        ["Priority", "Meaning", "Expected action"],
        [
            ["P0", "Production blocker", "Must be fixed before deployment."],
            ["P1", "Must address before scaling", "May not block an MVP but should be fixed before significant growth."],
            ["P2", "Should address during continued development", "Place in the near-term development backlog."],
            ["P3", "Acceptable MVP debt", "Temporarily acceptable; monitor and revisit."],
        ],
        widths=[0.75, 2.3, 3.65],
    )
    add_paragraph(
        doc,
        "The LLM assigns priority during detection and can revise it during validation. The backend normalizer maps common legacy severities and malformed variants into p0-p3 and adds display text, guidance, and UI color metadata.",
    )

    add_heading(doc, "8.2 Score and grade", 2)
    add_paragraph(
        doc,
        "The scoring engine is deterministic after validation. Priority weights are P0=4, P1=3, P2=2, and P3=1. Every issue contributes priority weight multiplied by its domain multiplier. Security has the highest multiplier because exploitable weaknesses can directly affect production safety.",
    )
    add_code_block(
        doc,
        """penalty = sum(priorityWeight(issue) * domainMultiplier(issue.domain))
score = max(0, 100 - penalty)

Grade thresholds:
A >= 90, B+ >= 80, B >= 70, C+ >= 60,
C >= 50, D >= 40, otherwise F""",
    )
    add_paragraph(
        doc,
        "This score is a product-level prioritization heuristic, not a calibrated probability of correctness, defect density, or formal security rating. Its purpose is to summarize validated findings consistently for comparison and triage.",
    )

    add_heading(doc, "8.3 Final issue representation", 2)
    add_paragraph(
        doc,
        "Each stored and returned issue contains its domain, file path, title, explanation, issue-specific suggestion, start and end lines, normalized priority, and optional rule and confidence fields. The response also constructs a GitHub file URL using the repository, default branch, path, and first affected line. Each domain group reports all, totalP0, totalP1, totalP2, totalP3, and items. Repository totals combine all domains.",
    )

    add_heading(doc, "9 Repair Prompt Generation and Regeneration", 1)
    add_paragraph(
        doc,
        "VibeCheck separates issue detection from remediation prompting. A validated issue already includes a suggestion, but the Generate Prompt feature creates a complete instruction that the user can paste into a coding assistant. The prompt includes repository and file context, line range, priority, issue title, explanation, suggested direction, behavioral constraints, requested code changes, rationale, and verification steps.",
    )
    add_code_block(
        doc,
        """You are a senior software engineering prompt designer.
Generate a repair prompt for a developer to paste into a coding assistant.

Return only JSON:
{
  "instructionAnalysis": "<what the user wants changed>",
  "fixPrompt": "<complete repair prompt>"
}

Context supplied to the model:
- repository name and URL
- issue domain, file URL, and line range
- P0-P3 priority and definition
- title, description, and suggested direction
- previous prompt, when one exists
- user regenerate instruction

The resulting prompt must request exact changes or a diff, explain why the
fix works, and request tests or manual checks unless the user opts out.""",
    )

    add_heading(doc, "9.1 First generation", 2)
    add_paragraph(
        doc,
        "The first request can use an empty body. Gemini creates version 1. If AI generation is unavailable, the backend can construct a deterministic local prompt from the issue record so the feature remains usable.",
    )

    add_heading(doc, "9.2 Regeneration with user intent", 2)
    add_paragraph(
        doc,
        "Once a prompt exists, regeneration requires regenerateInstruction, instruction, or promptInstruction. The instruction is normalized and limited to 1,000 characters. The AI receives the prior prompt, requested version, and user instruction and must return a meaningfully different prompt. It also returns instructionAnalysis, which records how the requested change was interpreted. Regeneration requires a successful AI response; provider failure or quota exhaustion is returned to the caller instead of pretending that the prompt was regenerated.",
    )
    add_paragraph(
        doc,
        "Example regenerate instruction: Make the prompt shorter, focus on a minimal code diff, preserve current behavior, and include one regression test for the affected path.",
        bold_lead="Example regenerate instruction:",
    )

    add_heading(doc, "10 Persistence Model and API Usage", 1)
    add_heading(doc, "10.1 MongoDB collections", 2)
    add_table(
        doc,
        ["Collection", "Stored data", "AI-related purpose"],
        [
            ["repositories", "Repository identity, branch, and last known size metadata.", "Avoids duplicating repository identity across runs."],
            ["roast_runs", "Check state, provider, model, score, grade, totals, timestamps, and errors.", "Represents one check and analysis lifecycle."],
            ["roast_issues", "Validated issue rows and generated prompt versions.", "Supports report detail and issue-level remediation."],
            ["analysis_file_caches", "Raw detections keyed by repository, Git blob SHA, domain, provider, model, and rules hash.", "Reduces repeated Phase 1 calls for unchanged files."],
            ["analysis_usages", "One record for each analysis execution counted against the daily quota.", "Limits expensive AI analyses per user and day."],
        ],
        widths=[1.45, 3.05, 2.2],
    )
    add_paragraph(
        doc,
        "The current schema does not persist the source contents or the personal OpenAI key. The file cache stores only raw structured detections. The current implementation also does not persist provider token usage, so historical token counts cannot be reconstructed exactly from MongoDB.",
    )

    add_heading(doc, "10.2 API workflow", 2)
    add_table(
        doc,
        ["Endpoint", "Purpose", "AI effect"],
        [
            ["POST /api/repos/check", "Validate and classify a repository; create roastId.", "No repository issue analysis."],
            ["POST /api/roasts/{roastId}/analyze", "Run six-domain analysis and return the completed report.", "Detection, validation, normalization, scoring, persistence."],
            ["GET /api/roasts/history", "List completed reports.", "Reads stored summaries; no new AI call."],
            ["GET /api/roasts/{roastId}", "Return a stored report and issues.", "Reads stored results; no new AI call."],
            ["POST /api/roasts/{roastId}/issues/{issueId}/prompt", "Generate or regenerate one repair prompt.", "Gemini prompt synthesis; regeneration analyzes user instruction."],
            ["DELETE /api/roasts/{roastId}", "Delete a completed report and its issue rows.", "No AI call."],
        ],
        widths=[2.85, 2.25, 1.6],
        font_size=8.2,
    )
    add_paragraph(
        doc,
        "An analysis execution, not the initial repository check, consumes the daily user quota. The default limit is ten analyses per user per Vietnam calendar day. The quota applies regardless of whether the request uses the system Gemini key or a personal OpenAI key.",
    )

    add_heading(doc, "11 Reproducibility Limitations and Reviewer Assessment", 1)
    add_heading(doc, "11.1 Reproducible elements", 2)
    add_bullets(
        doc,
        [
            "Repository size classification is deterministic for a given tree and ignore policy.",
            "File selection uses explicit scoring, byte limits, file limits, domain path hints, and a fixed Git commit SHA.",
            "Prompts are assembled from versioned templates, domain metadata, complete rule files, and explicit JSON schemas.",
            "Cache reuse is bound to file SHA, provider, model, domain, and detection-rule hash.",
            "Final normalization, P0-P3 mapping, score, grade, aggregation, and persistence are deterministic after model output.",
        ],
    )

    add_heading(doc, "11.2 Sources of variation", 2)
    add_bullets(
        doc,
        [
            "LLM output can vary even at low temperature, especially for borderline design issues.",
            "Model versions selected through environment variables can change behavior.",
            "A repository can change between the check and analyze calls because metadata is refreshed at analysis time.",
            "A system Gemini failure can produce heuristic fallback output that is not equivalent to a complete AI result.",
            "The validation prompt uses raw detections rather than the full source files, so the quality of Phase 2 depends on evidence captured in Phase 1.",
            "Priority-first limits mean that very large repositories are sampled rather than exhaustively verified.",
        ],
    )

    add_heading(doc, "11.3 Appropriate claims", 2)
    add_paragraph(
        doc,
        "VibeCheck should be presented as an AI-assisted repository review and production-debt triage system. It provides structured evidence, explanations, remediation suggestions, and repair prompts across recurring quality domains. It should not be described as a sound static analyzer, formal verifier, complete security scanner, or proof that a repository is production-safe. Human review and dedicated static-analysis tools remain necessary for high-stakes decisions.",
    )

    add_heading(doc, "11.4 Evidence that addresses the reviewer concern", 2)
    add_paragraph(
        doc,
        "The system's central AI capability can be assessed from four artifacts described in this report: the common Detection prompt, the common Validation prompt, twelve domain rule documents, and the repair-prompt generation template. Together they specify role, scope, definitions, evidence requirements, heuristics, priority semantics, JSON contracts, precision filtering, and remediation behavior. The prompts are not hidden application prose; they are built directly into each provider request and are therefore part of the implemented method.",
    )

    add_heading(doc, "Appendix A Prompt Construction Reference", 1)
    add_heading(doc, "A.1 Detection prompt fields", 2)
    add_table(
        doc,
        ["Prompt segment", "Value source", "Purpose"],
        [
            ["Role", "Domain configuration", "Frames expertise for maintainability, architecture, security, robustness, operational, or interaction analysis."],
            ["Task and recall policy", "Common template", "Requests all plausible issues and prohibits aggressive filtering."],
            ["Detection focus", "Domain configuration", "Narrows attention to the intended production-debt category."],
            ["Output schema", "Common template plus domain raw-name and metrics", "Forces structured evidence, lines, suggestion, priority, and heuristics."],
            ["Priority semantics", "Common template", "Defines P0-P3 in production terms."],
            ["Rule 1", "Domain Markdown file", "Supplies issue definitions, categories, thresholds, examples, and constraints."],
            ["Source batch", "Repository loader", "Supplies labeled file name, relative path, size, and content."],
        ],
        widths=[1.45, 1.8, 3.45],
    )

    add_heading(doc, "A.2 Validation prompt fields", 2)
    add_table(
        doc,
        ["Prompt segment", "Value source", "Purpose"],
        [
            ["Validation role", "Domain configuration", "Frames a precision-oriented expert."],
            ["Precision policy", "Common template", "Removes weak, speculative, incorrect, and duplicate issues."],
            ["Confidence guidance", "Common template", "Retains high or defensible medium findings and removes low confidence."],
            ["Suggestion requirements", "Common template", "Preserves or improves a concrete issue-specific remediation."],
            ["Rule 2", "Domain Markdown file", "Defines evidence strength, context, deduplication, and removal conditions."],
            ["Raw detections", "Cache plus current Phase 1 results", "Provides the candidate list to validate."],
        ],
        widths=[1.55, 1.85, 3.3],
    )

    add_heading(doc, "A.3 Domain rule files", 2)
    add_table(
        doc,
        ["Domain", "Detection Rule 1", "Validation Rule 2"],
        [
            ["Maintainability", "CODE_SMELL_RULE_1.md", "CODE_SMELL_RULE_2.md"],
            ["Architecture", "ARCHITECTURE_RULE_1.md", "ARCHITECTURE_RULE_2.md"],
            ["Security", "SECURITY_RULE_1.md", "SECURITY_RULE_2.md"],
            ["Robustness", "ROBUSTNESS_RULE_1.md", "ROBUSTNESS_RULE_2.md"],
            ["Operational", "OPERATIONAL_RULE_1.md", "OPERATIONAL_RULE_2.md"],
            ["Interaction", "INTERACTION_RULE_1.md", "INTERACTION_RULE_2.md"],
        ],
        widths=[1.45, 2.65, 2.6],
    )

    add_heading(doc, "Appendix B Implementation Map", 1)
    add_table(
        doc,
        ["Implementation area", "Primary source module"],
        [
            ["Repository metadata, ignore rules, priority loading", "src/services/githubRepo.service.js"],
            ["Six-domain orchestration, scoring, grade", "src/services/aiAnalysis.service.js"],
            ["Domain definitions and score multipliers", "src/services/analysis/analysisDomains.js"],
            ["Detection and validation prompt builders", "src/services/analysis/domainAnalysis.service.js"],
            ["Gemini and OpenAI HTTP clients", "src/services/analysis/llmClient.service.js"],
            ["Per-file raw detection cache", "src/services/analysis/fileAnalysisCache.service.js"],
            ["P0-P3 normalization and suggestions", "src/services/analysis/findingNormalizer.service.js"],
            ["Heuristic provider fallback", "src/services/analysis/localFallbackAnalysis.service.js"],
            ["Check, analyze, persistence, API report, repair prompt", "src/services/roast.service.js"],
            ["Issue, run, cache, repository, and usage schemas", "src/models/*.model.js"],
            ["OpenAPI endpoint contract", "src/routes/repo.routes.js and src/routes/roast.routes.js"],
        ],
        widths=[3.1, 3.6],
    )

    add_heading(doc, "Appendix C Suggested Paper Method Description", 1)
    add_paragraph(
        doc,
        "VibeCheck applies a two-stage LLM-assisted repository review. It first constructs a filtered and prioritized source-file set from a public GitHub repository. Six specialized domain analyzers then inspect the selected files for maintainability, architecture, security, robustness, operational readiness, and interaction reliability. For each domain, the Detection Phase combines source batches with a domain-specific Rule 1 prompt and emphasizes high recall, concrete evidence, heuristic metrics, file and line localization, and structured JSON. A separate Validation Phase combines the raw detections with Rule 2 and emphasizes high precision, realistic impact, evidence consistency, deduplication, and removal of low-confidence findings. The backend normalizes validated findings into a shared schema, classifies production priority from P0 to P3, computes a deterministic weighted score and grade, and stores issue-level results. A separate prompt-generation stage converts each validated issue into a reusable repair prompt and supports regeneration based on an explicit user instruction. File-level detection results are cached by Git blob SHA, provider, model, domain, and rule hash, while bounded parallel execution limits provider load.",
    )

    for section in doc.sections:
        add_footer(section)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
